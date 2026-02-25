"""
Exporter — produces Excel DD Grid.
Simplified version for vertical slice. Full version in Stage 9.
"""

import asyncio
import uuid

from src.core.logging import get_logger
from src.db.session import get_session
from src.queue.base_task import MASBaseTask
from src.queue.celery_app import celery_app
from src.storage.artifact_store import ArtifactPath, ArtifactStore


logger = get_logger(__name__)


try:
    import openpyxl
    from openpyxl.styles import PatternFill

    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False


SEVERITY_COLORS = {
    "critical": "FF0000",
    "high": "FF6600",
    "medium": "FFCC00",
    "low": "99CC00",
    "pass": "00CC00",
}

from src.db.models.result import CheckResult

async def _load_results(run_id: str) -> list["CheckResult"]:
    from src.db.repositories import ResultRepository

    async with get_session() as session:
        repo = ResultRepository(session)
        return await repo.get_by_run(uuid.UUID(run_id))


async def _save_export(run_id: str, export_ref: str | None) -> None:
    from src.db.models.export import Export
    from src.db.repositories import RunRepository

    async with get_session() as session:
        run = await RunRepository(session).get_by_id(uuid.UUID(run_id))
        if run:
            export = Export(run_id=run.id, ref=export_ref, format="xlsx")
            session.add(export)
            run.status = "completed"
            await session.flush()


async def _export_results_async(
    run_id: str,
    deal_id: str,
    tenant_id: str,
) -> dict:
    """
    Async core: loads check results, generates Excel, saves export artifact.
    """
    store = ArtifactStore()

    check_results = await _load_results(run_id)

    export_ref = None
    if HAS_OPENPYXL:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "DD Grid"
        ws.append(["Check ID", "Check Name", "Result", "Severity", "Value", "Confidence"])
        
        for cr in check_results:
            row = [
                cr.check_id,
                cr.check_name,
                cr.result,
                cr.severity,
                cr.value_text or "",
                f"{(cr.confidence or 0):.2f}",
            ]
            ws.append(row)
            severity = cr.severity.lower()
            color = SEVERITY_COLORS.get(severity, "FFFFFF")
            for cell in ws[ws.max_row]:
                cell.fill = PatternFill("solid", fgColor=color)
        
        import io

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        export_path = ArtifactPath.export(tenant_id, deal_id, run_id)
        export_ref = store.put_file(
            export_path + ".xlsx",
            buf.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    await _save_export(run_id, export_ref)
    
    logger.info("Export complete", run_id=run_id, rows=len(check_results))
    
    return {
        "run_id": run_id,
        "rows": len(check_results),
        "export_ref": export_ref,
    }


@celery_app.task(
    base=MASBaseTask,
    bind=True,
    name="workers.export_results",
    track_job=False,
    queue="default",
)
def export_results(
    self,
    run_id: str,
    deal_id: str,
    tenant_id: str,
) -> dict:
    """
    Celery entrypoint: sync wrapper around async export core.
    """
    logger.info("Starting export", run_id=run_id)

    return asyncio.run(
        _export_results_async(
            run_id=run_id,
            deal_id=deal_id,
            tenant_id=tenant_id,
        )
    )
