#!/usr/bin/env python3
"""
LegalMAS YAML → Excel exporter
Usage: python export_yaml_to_excel.py [--dir PATH] [--out OUTPUT.xlsx]
"""

import os, re, sys, yaml, argparse
from pathlib import Path
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

# ── CLI ──────────────────────────────────────────────────────────────────────

parser = argparse.ArgumentParser()
parser.add_argument('--dir', default='/home/vladi/projects/LegalMAS/playbooks/document_types')
parser.add_argument('--out', default='legalmAS_yaml_registry.xlsx')
args = parser.parse_args()

BASE_DIR = Path(args.dir)
OUT_FILE = args.out

# ── Palette ──────────────────────────────────────────────────────────────────

CLR_HEADER_BG  = '1F3A5F'   # тёмно-синий
CLR_HEADER_FG  = 'FFFFFF'
CLR_GROUP_BG   = 'D6E4F0'   # светло-голубой (group header)
CLR_GROUP_FG   = '1F3A5F'
CLR_ALT        = 'F0F6FB'   # чётные строки
CLR_WHITE      = 'FFFFFF'
CLR_ACCENT     = '2E75B6'   # акцент (ссылки, высокие)
CLR_WARN       = 'FFF2CC'   # жёлтый фон для scoring
CLR_HIGH_BG    = 'E2EFDA'   # зелёный фон high patterns
CLR_MED_BG     = 'FFF2CC'   # жёлтый фон medium
CLR_LOW_BG     = 'FCE4D6'   # оранжевый фон low

HDR_FONT   = Font(name='Calibri', bold=True, color=CLR_HEADER_FG, size=10)
HDR_FILL   = PatternFill('solid', fgColor=CLR_HEADER_BG)
HDR_ALIGN  = Alignment(horizontal='center', vertical='center', wrap_text=True)
BODY_FONT  = Font(name='Calibri', size=10)
BODY_ALIGN = Alignment(vertical='top', wrap_text=True, indent=1)
NUM_ALIGN  = Alignment(horizontal='right', vertical='top')
CTR_ALIGN  = Alignment(horizontal='center', vertical='top')
THIN       = Side(style='thin', color='BDD7EE')
BORDER     = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
TITLE_FONT = Font(name='Calibri', bold=True, size=14, color=CLR_HEADER_BG)
SUB_FONT   = Font(name='Calibri', size=10, color='595959', italic=True)

# ── Helpers ──────────────────────────────────────────────────────────────────

def cell_set(ws, row, col, value, font=None, fill=None, alignment=None, border=None, num_format=None):
    c = ws.cell(row=row, column=col, value=value)
    if font:      c.font = font
    if fill:      c.fill = fill
    if alignment: c.alignment = alignment
    if border:    c.border = border
    if num_format: c.number_format = num_format
    return c

def keywords_summary(kw):
    if not isinstance(kw, dict): return ''
    parts = []
    for lvl in ('high', 'medium', 'low'):
        n = len(kw.get(lvl, []))
        if n: parts.append(f'{lvl.upper()}: {n}')
    return ' | '.join(parts)

def excl_summary(excl):
    if not isinstance(excl, list): return ''
    h = sum(1 for x in excl if isinstance(x, dict) and x.get('severity') == 'hard')
    s = sum(1 for x in excl if isinstance(x, dict) and x.get('severity') == 'soft')
    p = sum(1 for x in excl if isinstance(x, str))
    parts = []
    if h: parts.append(f'hard: {h}')
    if s: parts.append(f'soft: {s}')
    if p: parts.append(f'plain: {p}')
    return ' | '.join(parts)

def pattern_lines(kw, level):
    items = kw.get(level, []) if isinstance(kw, dict) else []
    out = []
    for item in items:
        if isinstance(item, str):
            out.append(item)
        elif isinstance(item, dict):
            p   = item.get('pattern', '')
            w   = item.get('weight', '')
            d   = item.get('distinctiveness', '')
            ctx = item.get('context', '')
            desc= item.get('description', '')
            tag = f'[w={w} d={d} ctx={ctx}]'
            out.append(f'{tag} {p}' + (f'  — {desc}' if desc else ''))
    return '\n'.join(out)

def excl_lines(excl):
    if not isinstance(excl, list): return ''
    out = []
    for item in excl:
        if isinstance(item, str):
            out.append(f'[plain] {item}')
        elif isinstance(item, dict):
            p   = item.get('pattern', '')
            sev = item.get('severity', '')
            w   = item.get('weight', '')
            desc= item.get('description', '')
            out.append(f'[{sev} w={w}] {p}' + (f'  — {desc}' if desc else ''))
    return '\n'.join(out)

# ── Scan YAMLs ────────────────────────────────────────────────────────────────

yaml_files = sorted(BASE_DIR.rglob('*.yaml'))
records = []
for fpath in yaml_files:
    if fpath.name == 'registry.yaml':
        continue
    try:
        data = yaml.safe_load(fpath.read_text(encoding='utf-8'))
    except Exception as e:
        print(f'SKIP {fpath.name}: {e}')
        continue
    if not isinstance(data, dict):
        continue

    meta    = data.get('metadata', {}) or {}
    engine  = data.get('engine_expectations', {}) or {}
    cr      = data.get('classification_runtime_v1', {}) or {}
    scoring = cr.get('scoring', {}) or {}
    rules   = cr.get('rules', {}) or {}
    kw      = rules.get('keywords', {}) or {}
    excl    = rules.get('exclusions', []) or []
    fn_pat  = rules.get('filename_patterns', []) or []
    fo_pat  = rules.get('folder_patterns', []) or []
    trig    = data.get('trigger_definitions', {}) or {}
    shared  = data.get('shared_patterns', {}) or {}

    records.append({
        'file':           str(fpath.relative_to(BASE_DIR)),
        'group':          fpath.parent.name,
        'type_code':      data.get('type', ''),
        'name':           data.get('name', ''),
        'version':        str(data.get('version', '')),
        'parent':         data.get('parent', ''),
        'description':    (data.get('description', '') or '').strip(),
        'category':       meta.get('category', ''),
        'subcategory':    meta.get('subcategory', ''),
        'complexity':     meta.get('complexity', ''),
        'risk_level':     meta.get('risk_level', ''),
        'jurisdiction':   meta.get('jurisdiction', ''),
        'language':       meta.get('language', ''),
        'legal_base':     meta.get('legal_base', ''),
        'doc_types_supported': '\n'.join(meta.get('doc_types_supported', [])),
        'evidence_style':      engine.get('evidence_style', ''),
        'output_style':        engine.get('output_style', ''),
        'accept_threshold':    scoring.get('accept_threshold', ''),
        'strong_signal_thr':   scoring.get('strong_signal_threshold', ''),
        'ambiguity_margin':    scoring.get('ambiguity_margin', ''),
        'iq_weight':           scoring.get('input_quality_weight', ''),
        'kw_summary':          keywords_summary(kw),
        'kw_high_count':       len(kw.get('high', [])),
        'kw_med_count':        len(kw.get('medium', [])),
        'kw_low_count':        len(kw.get('low', [])),
        'excl_summary':        excl_summary(excl),
        'excl_count':          len(excl),
        'fn_patterns':         len(fn_pat),
        'folder_patterns':     len(fo_pat),
        'triggers_count':      len(trig),
        'shared_count':        len(shared),
        # detail
        '_kw_high':  pattern_lines(kw, 'high'),
        '_kw_med':   pattern_lines(kw, 'medium'),
        '_kw_low':   pattern_lines(kw, 'low'),
        '_excl':     excl_lines(excl),
        '_fn_pat':   '\n'.join(fn_pat),
        '_fo_pat':   '\n'.join(fo_pat),
    })

print(f'Parsed {len(records)} YAML files')

# ── Workbook ──────────────────────────────────────────────────────────────────

wb = Workbook()

# ═══════════════════════════════════════════════════════════════════
# Лист 1: Overview
# ═══════════════════════════════════════════════════════════════════

ws1 = wb.active
ws1.title = 'Overview'
ws1.sheet_view.showGridLines = False
ws1.column_dimensions['A'].width = 3

# Title
ws1.merge_cells('B2:T2')
ws1['B2'] = 'LegalMAS — Реестр инструкций типов документов'
ws1['B2'].font  = TITLE_FONT
ws1['B2'].alignment = Alignment(horizontal='left', vertical='center')
ws1.row_dimensions[2].height = 30

ws1.merge_cells('B3:T3')
ws1['B3'] = f'Сформировано: {datetime.now().strftime("%Y-%m-%d %H:%M")}  |  Файлов: {len(records)}'
ws1['B3'].font = SUB_FONT
ws1['B3'].alignment = Alignment(horizontal='left', vertical='center')
ws1.row_dimensions[3].height = 18

# Header row
OVR_COLS = [
    ('Группа',            14),
    ('Type Code',         28),
    ('Name',              30),
    ('Ver',                5),
    ('Parent',            22),
    ('Complexity',        11),
    ('Risk',              10),
    ('Legal Base',        28),
    ('accept_thr',        10),
    ('strong_thr',        10),
    ('ambiguity',          9),
    ('KW High',            8),
    ('KW Med',             8),
    ('KW Low',             8),
    ('Excl',               8),
    ('FN Pat',             7),
    ('Triggers',           8),
    ('KW Summary',        22),
    ('Excl Summary',      22),
]

HR = 5
for ci, (hdr, w) in enumerate(OVR_COLS, start=2):
    c = ws1.cell(row=HR, column=ci, value=hdr)
    c.font      = HDR_FONT
    c.fill      = HDR_FILL
    c.alignment = HDR_ALIGN
    c.border    = BORDER
    ws1.column_dimensions[get_column_letter(ci)].width = w
ws1.row_dimensions[HR].height = 32
ws1.freeze_panes = f'B{HR+1}'

# Data rows
ALT_FILL  = PatternFill('solid', fgColor=CLR_ALT)
HIGH_FILL = PatternFill('solid', fgColor='E2EFDA')

for ri, rec in enumerate(records):
    row = HR + 1 + ri
    fill = ALT_FILL if ri % 2 == 0 else PatternFill('solid', fgColor=CLR_WHITE)
    vals = [
        rec['group'],
        rec['type_code'],
        rec['name'],
        rec['version'],
        rec['parent'],
        rec['complexity'],
        rec['risk_level'],
        rec['legal_base'],
        rec['accept_threshold'],
        rec['strong_signal_thr'],
        rec['ambiguity_margin'],
        rec['kw_high_count'],
        rec['kw_med_count'],
        rec['kw_low_count'],
        rec['excl_count'],
        rec['fn_patterns'],
        rec['triggers_count'],
        rec['kw_summary'],
        rec['excl_summary'],
    ]
    for ci, val in enumerate(vals, start=2):
        c = ws1.cell(row=row, column=ci, value=val)
        c.fill   = fill
        c.border = BORDER
        is_num = isinstance(val, (int, float))
        c.font      = BODY_FONT
        c.alignment = NUM_ALIGN if is_num else (CTR_ALIGN if ci in (2,6,7,9,10,11,12,13,14,15,16,17) else BODY_ALIGN)
    ws1.row_dimensions[row].height = 20

# Excel Table
last_row  = HR + len(records)
last_col  = HR + len(OVR_COLS) - 4  # col index
tbl_ref   = f'B{HR}:{get_column_letter(len(OVR_COLS)+1)}{last_row}'
tbl = Table(displayName='OverviewTable', ref=tbl_ref)
tbl.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
ws1.add_table(tbl)


# ═══════════════════════════════════════════════════════════════════
# Лист 2: Descriptions
# ═══════════════════════════════════════════════════════════════════

ws2 = wb.create_sheet('Descriptions')
ws2.sheet_view.showGridLines = False
ws2.column_dimensions['A'].width = 3

DESC_COLS = [
    ('Type Code',    28),
    ('Name',         32),
    ('Description',  70),
    ('doc_types_supported', 40),
]
for ci, (h, w) in enumerate(DESC_COLS, start=2):
    c = ws2.cell(row=2, column=ci, value=h)
    c.font = HDR_FONT; c.fill = HDR_FILL; c.alignment = HDR_ALIGN; c.border = BORDER
    ws2.column_dimensions[get_column_letter(ci)].width = w
ws2.row_dimensions[2].height = 28
ws2.freeze_panes = 'B3'

for ri, rec in enumerate(records):
    row  = 3 + ri
    fill = ALT_FILL if ri % 2 == 0 else PatternFill('solid', fgColor=CLR_WHITE)
    for ci, val in enumerate([
        rec['type_code'], rec['name'],
        rec['description'], rec['doc_types_supported']
    ], start=2):
        c = ws2.cell(row=row, column=ci, value=val)
        c.font = BODY_FONT; c.fill = fill; c.border = BORDER
        c.alignment = Alignment(vertical='top', wrap_text=True, indent=1)
    ws2.row_dimensions[row].height = max(15, min(90, rec['description'].count('\n') * 13 + 15))


# ═══════════════════════════════════════════════════════════════════
# Лист 3: Keywords & Exclusions (detail)
# ═══════════════════════════════════════════════════════════════════

ws3 = wb.create_sheet('Keywords & Exclusions')
ws3.sheet_view.showGridLines = False
ws3.column_dimensions['A'].width = 3

KW_COLS = [
    ('Type Code',     28),
    ('HIGH patterns', 70),
    ('MEDIUM patterns',70),
    ('LOW patterns',  60),
    ('Exclusions',    70),
]
HIGH_HDR_FILL = PatternFill('solid', fgColor='375623')
MED_HDR_FILL  = PatternFill('solid', fgColor='7F6000')
LOW_HDR_FILL  = PatternFill('solid', fgColor='843C0C')
EXCL_HDR_FILL = PatternFill('solid', fgColor='1F3A5F')

hdr_fills = [HDR_FILL, HIGH_HDR_FILL, MED_HDR_FILL, LOW_HDR_FILL, EXCL_HDR_FILL]

for ci, ((h, w), hf) in enumerate(zip(KW_COLS, hdr_fills), start=2):
    c = ws3.cell(row=2, column=ci, value=h)
    c.font = HDR_FONT; c.fill = hf; c.alignment = HDR_ALIGN; c.border = BORDER
    ws3.column_dimensions[get_column_letter(ci)].width = w
ws3.row_dimensions[2].height = 28
ws3.freeze_panes = 'B3'

HIGH_ROW_FILL = PatternFill('solid', fgColor='EBF5EB')
MED_ROW_FILL  = PatternFill('solid', fgColor='FFFAEB')
LOW_ROW_FILL  = PatternFill('solid', fgColor='FEF4EE')
EXCL_ROW_FILL = PatternFill('solid', fgColor='EFF3F8')

for ri, rec in enumerate(records):
    row = 3 + ri
    fills_per_col = [ALT_FILL if ri%2==0 else PatternFill('solid', fgColor=CLR_WHITE),
                     HIGH_ROW_FILL, MED_ROW_FILL, LOW_ROW_FILL, EXCL_ROW_FILL]
    vals = [rec['type_code'], rec['_kw_high'], rec['_kw_med'], rec['_kw_low'], rec['_excl']]
    for ci, (val, rf) in enumerate(zip(vals, fills_per_col), start=2):
        c = ws3.cell(row=row, column=ci, value=val)
        c.font = Font(name='Consolas', size=9) if ci > 2 else BODY_FONT
        c.fill = rf; c.border = BORDER
        c.alignment = Alignment(vertical='top', wrap_text=True, indent=1)
    # row height based on line count
    lines = max(rec['_kw_high'].count('\n'), rec['_kw_med'].count('\n'),
                rec['_kw_low'].count('\n'), rec['_excl'].count('\n')) + 2
    ws3.row_dimensions[row].height = min(lines * 14, 300)


# ═══════════════════════════════════════════════════════════════════
# Лист 4: Scoring & Patterns
# ═══════════════════════════════════════════════════════════════════

ws4 = wb.create_sheet('Scoring & File Patterns')
ws4.sheet_view.showGridLines = False
ws4.column_dimensions['A'].width = 3

SC_COLS = [
    ('Type Code',          28),
    ('accept_threshold',   15),
    ('strong_signal_thr',  15),
    ('ambiguity_margin',   15),
    ('iq_weight',          12),
    ('Filename Patterns',  55),
    ('Folder Patterns',    40),
]
for ci, (h, w) in enumerate(SC_COLS, start=2):
    c = ws4.cell(row=2, column=ci, value=h)
    c.font = HDR_FONT; c.fill = HDR_FILL; c.alignment = HDR_ALIGN; c.border = BORDER
    ws4.column_dimensions[get_column_letter(ci)].width = w
ws4.row_dimensions[2].height = 28
ws4.freeze_panes = 'B3'

WARN_FILL = PatternFill('solid', fgColor='FFF2CC')

for ri, rec in enumerate(records):
    row  = 3 + ri
    base = ALT_FILL if ri % 2 == 0 else PatternFill('solid', fgColor=CLR_WHITE)
    vals = [
        rec['type_code'],
        rec['accept_threshold'],
        rec['strong_signal_thr'],
        rec['ambiguity_margin'],
        rec['iq_weight'],
        rec['_fn_pat'],
        rec['_fo_pat'],
    ]
    for ci, val in enumerate(vals, start=2):
        c = ws4.cell(row=row, column=ci, value=val)
        c.border = BORDER
        no_scoring = ci in (3,4,5,6) and val == ''
        c.fill = WARN_FILL if no_scoring else base
        c.font = Font(name='Consolas', size=9) if ci >= 7 else BODY_FONT
        c.alignment = (NUM_ALIGN if ci in (3,4,5,6) else Alignment(vertical='top', wrap_text=True, indent=1))
    fn_lines = rec['_fn_pat'].count('\n') + 2
    ws4.row_dimensions[row].height = max(15, min(fn_lines * 13, 120))


# ═══════════════════════════════════════════════════════════════════
# Лист 5: Triggers & Shared
# ═══════════════════════════════════════════════════════════════════

ws5 = wb.create_sheet('Triggers & Shared')
ws5.sheet_view.showGridLines = False
ws5.column_dimensions['A'].width = 3

TR_COLS = [
    ('Type Code',      28),
    ('Triggers count', 14),
    ('Shared count',   13),
    ('Trigger names',  55),
    ('Shared keys',    40),
]
for ci, (h, w) in enumerate(TR_COLS, start=2):
    c = ws5.cell(row=2, column=ci, value=h)
    c.font = HDR_FONT; c.fill = HDR_FILL; c.alignment = HDR_ALIGN; c.border = BORDER
    ws5.column_dimensions[get_column_letter(ci)].width = w
ws5.row_dimensions[2].height = 28
ws5.freeze_panes = 'B3'

for ri, rec in enumerate(records):
    row  = 3 + ri
    fill = ALT_FILL if ri % 2 == 0 else PatternFill('solid', fgColor=CLR_WHITE)

    # Re-read YAML for trigger names & shared keys (need raw data)
    yaml_path = BASE_DIR / rec['file']
    try:
        raw = yaml.safe_load(yaml_path.read_text(encoding='utf-8'))
        trig_names  = '\n'.join((raw.get('trigger_definitions', {}) or {}).keys())
        shared_keys = '\n'.join((raw.get('shared_patterns', {}) or {}).keys())
    except Exception:
        trig_names = ''; shared_keys = ''

    vals = [rec['type_code'], rec['triggers_count'], rec['shared_count'], trig_names, shared_keys]
    for ci, val in enumerate(vals, start=2):
        c = ws5.cell(row=row, column=ci, value=val)
        c.font = BODY_FONT; c.fill = fill; c.border = BORDER
        c.alignment = NUM_ALIGN if ci in (3,4) else Alignment(vertical='top', wrap_text=True, indent=1)
    lines = trig_names.count('\n') + 2
    ws5.row_dimensions[row].height = max(15, min(lines * 13, 200))


# ── Save ─────────────────────────────────────────────────────────────────────

wb.save(OUT_FILE)
print(f'Saved: {OUT_FILE}')

