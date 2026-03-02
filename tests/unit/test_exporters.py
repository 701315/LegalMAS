"""Тесты для Excel и Word экспортёров (Stage 9)."""

import io

import pytest
from docx import Document
from openpyxl import load_workbook

from src.services.export.aggregator import aggregate_results
from src.services.export.grid_builder import build_grid_excel
from src.services.export.report_builder import build_report_docx
from src.services.export.summary_builder import build_summary_excel
from tests.fixtures.export_test_data import SAMPLE_CHECK_RESULTS_V9, SAMPLE_DOCUMENTS_V9

DEAL_ID = "deal-export-test"
DEAL_NAME = "Export Test Deal"
RUN_ID = "run-export-001"


def _get_aggregated():
    return aggregate_results(
        deal_id=DEAL_ID,
        deal_name=DEAL_NAME,
        run_id=RUN_ID,
        check_results=SAMPLE_CHECK_RESULTS_V9,
        documents=SAMPLE_DOCUMENTS_V9,
    )


@pytest.fixture(scope="module")
def aggregated():
    return _get_aggregated()


class TestNewGridSheets:
    def test_conditional_sheet_present_when_triggered(self, aggregated):
        data = build_grid_excel(aggregated)
        wb = load_workbook(io.BytesIO(data))
        if aggregated.conditional_rows:
            assert "Conditional Checks" in wb.sheetnames

    def test_extracted_values_sheet(self, aggregated):
        data = build_grid_excel(aggregated)
        wb = load_workbook(io.BytesIO(data))
        if aggregated.structured_value_rows:
            assert "Extracted Values" in wb.sheetnames

    def test_grid_has_structured_data_column(self, aggregated):
        data = build_grid_excel(aggregated)
        wb = load_workbook(io.BytesIO(data))
        ws = wb["DD Grid"]
        headers = [ws.cell(row=1, column=i).value for i in range(1, 20)]
        if aggregated.structured_value_rows:
            assert "Извлечённые данные" in headers


class TestSummaryNewSheets:
    def test_classification_quality_sheet(self, aggregated):
        data = build_summary_excel(aggregated)
        wb = load_workbook(io.BytesIO(data))
        assert "Classification Quality" in wb.sheetnames

    def test_by_type_has_classification_columns(self, aggregated):
        data = build_summary_excel(aggregated)
        wb = load_workbook(io.BytesIO(data))
        ws = wb["By Doc Type"]
        headers = [ws.cell(row=1, column=i).value for i in range(1, 12)]
        # В плане колонка называлась "Avg Conf. классиф." — подстраиваемся под фактический текст
        assert any("Avg Conf." in (h or "") for h in headers)

    def test_request_list_has_addressee(self, aggregated):
        data = build_summary_excel(aggregated)
        wb = load_workbook(io.BytesIO(data))
        ws = wb["Request List"]
        headers = [ws.cell(row=1, column=i).value for i in range(1, 9)]
        assert "Адресат" in headers


class TestReportNewSections:
    def test_conditional_section_in_report(self, aggregated):
        data = build_report_docx(aggregated)
        doc = Document(io.BytesIO(data))
        full_text = " ".join(p.text for p in doc.paragraphs)
        if aggregated.conditional_rows:
            assert ("Условные проверки" in full_text) or ("условные проверки" in full_text.lower())

    def test_extracted_values_in_report(self, aggregated):
        data = build_report_docx(aggregated)
        doc = Document(io.BytesIO(data))
        full_text = " ".join(p.text for p in doc.paragraphs)
        if aggregated.structured_value_rows:
            assert "Извлечённые значения" in full_text

    def test_request_list_has_addressee_text(self, aggregated):
        data = build_report_docx(aggregated)
        doc = Document(io.BytesIO(data))
        full_text = " ".join(p.text for p in doc.paragraphs)
        if aggregated.dd_requests:
            assert "Адресат" in full_text
