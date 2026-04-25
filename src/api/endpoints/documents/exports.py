# src/api/endpoints/documents/exports.py

"""
Document exports API.

Содержит endpoints для экспорта списка документов сделки в:
- Excel
- Word
- PDF
"""

from __future__ import annotations

import logging
from io import BytesIO
from pathlib import Path
from typing import Optional
from uuid import UUID

from fastapi import APIRouter, Depends, HTTPException, status
from fastapi.responses import Response
from fpdf import FPDF
from pydantic import BaseModel, Field
from sqlalchemy.ext.asyncio import AsyncSession

from api.deps.auth import CurrentUser, get_current_user
from api.deps.security import require_project_role
from core.document_columns import (
    DocumentColumnDef,
    DocumentColumnId,
    get_columns_by_ids,
)
from db.models.document import Document
from db.repositories.document_repo import DocumentRepository
from db.session import get_db
from services.document_composition_query import DocumentCompositionQueryService

logger = logging.getLogger(__name__)

router = APIRouter()

BASE_DIR = Path(__file__).resolve().parent.parent.parent.parent
FONT_PATH = BASE_DIR / "resources" / "fonts" / "DejaVuSans.ttf"


class DocumentsExportRequest(BaseModel):
    """
    Параметры экспорта списка документов.

    visible_column_ids:
        Список ID колонок, которые нужно включить в экспорт.
        Если None или пусто — используются все доступные колонки.

    deal_id:
        ID сделки, для которой выгружается список документов.
    """

    deal_id: UUID
    visible_column_ids: Optional[list[DocumentColumnId]] = Field(default=None)


def build_table_rows(
    documents: list[Document],
    visible_column_ids: Optional[list[DocumentColumnId]],
    composition_by_doc_id: dict[UUID, object] | None = None,
) -> tuple[list[str], list[list[str]]]:
    """
    Построить заголовки и строки таблицы для экспорта на основе списка документов
    и выбранных колонок.

    Для непрофессионального программиста:
    - documents — это строки из таблицы documents (ORM-модели Document);
    - composition_by_doc_id — словарь с краткой информацией о сегментах
      (single/multi, сколько сегментов, основной сегмент) для каждого документа;
    - visible_column_ids определяет, какие именно колонки попадут в экспорт.
    """
    column_defs: list[DocumentColumnDef] = get_columns_by_ids(visible_column_ids)
    headers = [column.label for column in column_defs]

    rows: list[list[str]] = []
    composition_by_doc_id = composition_by_doc_id or {}

    for doc in documents:
        row: list[str] = []

        for column in column_defs:
            value = _get_export_value(
                doc=doc,
                column=column,
                composition=composition_by_doc_id.get(doc.id),
            )

            if isinstance(value, bool):
                row.append("Да" if value else "Нет")
            elif value is None:
                row.append("")
            else:
                row.append(str(value))

        rows.append(row)

    return headers, rows


def _get_export_value(
    doc: Document,
    column: DocumentColumnDef,
    composition: object | None = None,
) -> object:
    """
    Вернуть значение колонки документа с учётом вычисляемых полей и нормализации.

    Для непрофессионального программиста:
    - column.field может быть:
        * именем реального поля ORM-модели Document (например, "filename");
        * псевдонимом вычисляемого поля (например, "applicable_triggers_count"
          или composition-поля).
    - composition — это объект DocumentCompositionSummary или None,
      который мы заранее посчитали в отдельном сервисе.
    """
    field = column.field

    # Ветвь для composition-полей (они не лежат в модели Document непосредственно).
    if field in {
        "composition_type",
        "segment_count",
        "primary_segment_label",
        "is_composite_file",
    }:
        if composition is None:
            # Сегментация ещё не выполнялась или работает в legacy-режиме.
            if field == "segment_count":
                return 0
            if field == "is_composite_file":
                return False
            return None

        # composition — это Pydantic-модель DocumentCompositionSummary
        # или другой объект с одноимёнными атрибутами.
        return getattr(composition, field, None)

    # Базовый случай: читаем реальное поле Document.
    value = getattr(doc, field, None)

    if field == "applicable_triggers_count":
        return len(getattr(doc, "applicable_triggers", []) or [])

    if field == "status":
        return doc.status or ""

    if field == "risk_level":
        return doc.risk_level or ""

    return value


async def _load_documents_for_export(
    db: AsyncSession,
    user: CurrentUser,
    deal_id: UUID,
) -> tuple[list[Document], dict[UUID, object]]:
    """
    Проверить права и загрузить документы сделки для экспорта,
    а также заранее построить composition для каждого документа.

    Для непрофессионального программиста:
    - этот helper возвращает сразу два объекта:
        * список документов;
        * словарь {document_id: composition_summary}.
    - так мы избегаем N+1 запросов и держим всю бизнес-логику экспорта
      (включая composition) в одном месте.
    """
    await require_project_role(
        db=db,
        user=user,
        deal_id=deal_id,
        allowed_roles=("viewer", "editor", "admin"),
    )

    doc_repo = DocumentRepository(db)
    documents = await doc_repo.list_by_deal_id(deal_id)

    if not documents:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Документы не найдены",
        )

    # Строим composition для всех документов одним batch-запросом.
    composition_service = DocumentCompositionQueryService(db)
    composition_by_doc_id = await composition_service.build_for_documents(documents)

    return documents, composition_by_doc_id


def _build_download_headers(filename: str) -> dict[str, str]:
    """
    Построить заголовки ответа для скачивания файла.
    """
    return {
        "Content-Disposition": f'attachment; filename="{filename}"',
        "Access-Control-Expose-Headers": "Content-Disposition",
    }


@router.post(
    "/export/excel",
    summary="Экспорт списка документов сделки в Excel",
)
async def export_documents_excel(
    body: DocumentsExportRequest,
    db: AsyncSession = Depends(get_db),
    user: CurrentUser = Depends(get_current_user),
):
    """
    Экспортирует список документов сделки в Excel.

    - Проверяет права на сделку.
    - Собирает таблицу по выбранным колонкам (включая composition-поля).
    - Возвращает xlsx-файл.
    """
    documents, composition_by_doc_id = await _load_documents_for_export(
        db=db,
        user=user,
        deal_id=body.deal_id,
    )
    headers, rows = build_table_rows(
        documents=documents,
        visible_column_ids=body.visible_column_ids,
        composition_by_doc_id=composition_by_doc_id,
    )

    try:
        import openpyxl

        workbook = openpyxl.Workbook()
        worksheet = workbook.active
        worksheet.title = "Documents"

        worksheet.append(headers)
        for row in rows:
            worksheet.append(row)

        stream = BytesIO()
        workbook.save(stream)
        stream.seek(0)
    except HTTPException:
        raise
    except Exception as exc:
        logger.exception(
            "Failed to export documents to Excel",
            extra={
                "deal_id": str(body.deal_id),
                "user_id": str(user.id),
                "error": str(exc),
            },
        )
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail="Excel export temporarily unavailable",
        )

    filename = f"deal_{body.deal_id}_documents.xlsx"
    return Response(
        content=stream.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=_build_download_headers(filename),
    )


@router.post(
    "/export/word",
    summary="Экспорт списка документов сделки в Word (docx)",
)
async def export_documents_word(
    body: DocumentsExportRequest,
    db: AsyncSession = Depends(get_db),
    user: CurrentUser = Depends(get_current_user),
):
    """
    Экспортирует список документов сделки в Word (docx).
    """
    documents, composition_by_doc_id = await _load_documents_for_export(
        db=db,
        user=user,
        deal_id=body.deal_id,
    )
    headers, rows = build_table_rows(
        documents=documents,
        visible_column_ids=body.visible_column_ids,
        composition_by_doc_id=composition_by_doc_id,
    )

    try:
        from docx import Document as DocxDocument

        docx = DocxDocument()
        table = docx.add_table(rows=1, cols=len(headers))

        header_cells = table.rows[0].cells
        for index, header in enumerate(headers):
            header_cells[index].text = header

        for row in rows:
            row_cells = table.add_row().cells
            for index, cell_value in enumerate(row):
                row_cells[index].text = cell_value

        stream = BytesIO()
        docx.save(stream)
        stream.seek(0)
    except HTTPException:
        raise
    except Exception as exc:
        logger.exception(
            "Failed to export documents to Word",
            extra={
                "deal_id": str(body.deal_id),
                "user_id": str(user.id),
                "error": str(exc),
            },
        )
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail="Word export temporarily unavailable",
        )

    filename = f"deal_{body.deal_id}_documents.docx"
    return Response(
        content=stream.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers=_build_download_headers(filename),
    )


@router.post(
    "/export/pdf",
    summary="Экспорт списка документов сделки в PDF",
)
async def export_documents_pdf(
    body: DocumentsExportRequest,
    db: AsyncSession = Depends(get_db),
    user: CurrentUser = Depends(get_current_user),
):
    """
    Экспортирует список документов сделки в PDF.
    """
    documents, composition_by_doc_id = await _load_documents_for_export(
        db=db,
        user=user,
        deal_id=body.deal_id,
    )
    headers, rows = build_table_rows(
        documents=documents,
        visible_column_ids=body.visible_column_ids,
        composition_by_doc_id=composition_by_doc_id,
    )

    if not FONT_PATH.exists():
        logger.error(
            "PDF export font file not found",
            extra={"font_path": str(FONT_PATH)},
        )
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail="PDF export temporarily unavailable",
        )

    try:
        pdf = FPDF()
        pdf.add_page()
        pdf.add_font("DejaVu", "", FONT_PATH.as_posix(), uni=True)
        pdf.set_font("DejaVu", size=10)

        line_height = pdf.font_size * 1.5
        col_width = pdf.epw / max(len(headers), 1)

        for header in headers:
            pdf.cell(col_width, line_height, txt=header[:50], border=1)
        pdf.ln(line_height)

        for row in rows:
            for cell_value in row:
                pdf.cell(col_width, line_height, txt=str(cell_value)[:50], border=1)
            pdf.ln(line_height)

        out_bytes = bytes(pdf.output(dest="S"))
    except HTTPException:
        raise
    except Exception as exc:
        logger.exception(
            "Failed to export documents to PDF",
            extra={
                "deal_id": str(body.deal_id),
                "user_id": str(user.id),
                "error": str(exc),
            },
        )
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail="PDF export temporarily unavailable",
        )

    filename = f"deal_{body.deal_id}_documents.pdf"
    return Response(
        content=out_bytes,
        media_type="application/pdf",
        headers=_build_download_headers(filename),
    )

    