"""
Exporter — produces Excel DD Grid.
Simplified version for vertical slice. Full version in Stage 9.
"""

import io
import uuid

from sqlalchemy.ext.asyncio import AsyncSession

from src.core.logging import get_logger
from src.db.models.export import Export
from src.db.models.result import CheckResult
from src.db.repositories import ResultRepository, RunRepository
from src.db.uow import db_uow
from src.queue.base_task import MASBaseTask
from src.queue.celery_app import celery_app
from src.storage.artifact_store import ArtifactPath, ArtifactStore

logger = get_logger(__name__)


try:
    import openpyxl
    from openpyxl.styles import PatternFill

    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False


SEVERITY_COLORS = {
    "critical": "FF0000",
    "high": "FF6600",
    "medium": "FFCC00",
    "low": "99CC00",
    "pass": "00CC00",
}


async def _load_results(run_id: str, db: AsyncSession) -> list["CheckResult"]:
    repo = ResultRepository(db)
    return await repo.get_by_run(uuid.UUID(run_id))


async def _save_export(run_id: str, export_ref: str | None, db: AsyncSession) -> None:
    run_repo = RunRepository(db)
    run = await run_repo.get_by_id(uuid.UUID(run_id))
    if run:
        export = Export(run_id=run.id, ref=export_ref, format="xlsx")
        db.add(export)
        run.status = "completed"
        await db.flush()


async def _export_results_async(
    run_id: str,
    deal_id: str,
    tenant_id: str,
    db: AsyncSession,
) -> dict:
    """
    Async core: loads check results, generates Excel, saves export artifact.

    Все операции с БД выполняются через переданную AsyncSession (db),
    чтобы одна задача worker'а работала в рамках одного unit of work.
    """
    store = ArtifactStore()

    check_results = await _load_results(run_id, db)

    export_ref = None
    if HAS_OPENPYXL:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "DD Grid"
        ws.append(["Check ID", "Check Name", "Result", "Severity", "Value", "Confidence"])

        for cr in check_results:
            row = [
                cr.check_id,
                cr.check_name,
                cr.result,
                cr.severity,
                cr.value_text or "",
                f"{(cr.confidence or 0):.2f}",
            ]
            ws.append(row)
            severity = (cr.severity or "").lower()
            color = SEVERITY_COLORS.get(severity, "FFFFFF")
            for cell in ws[ws.max_row]:
                cell.fill = PatternFill("solid", fgColor=color)

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        export_path = ArtifactPath.export(tenant_id, deal_id, run_id)
        export_ref = store.put_file(
            export_path + ".xlsx",
            buf.getvalue(),
            content_type=("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
        )

    await _save_export(run_id, export_ref, db)

    logger.info("Export complete", run_id=run_id, rows=len(check_results))

    return {
        "run_id": run_id,
        "rows": len(check_results),
        "export_ref": export_ref,
    }


@celery_app.task(
    base=MASBaseTask,
    bind=True,
    name="workers.export_results",
    track_job=False,
    queue="default",
)
async def export_results(
    self,
    run_id: str,
    deal_id: str,
    tenant_id: str,
) -> dict:
    """
    Celery entrypoint: async wrapper around export core.

    Для каждой таски создаётся отдельный UnitOfWork/AsyncSession,
    внутри которого выполняются все операции с БД для данного run_id.
    """
    logger.info("Starting export", run_id=run_id)

    async with db_uow() as uow:
        return await _export_results_async(
            run_id=run_id,
            deal_id=deal_id,
            tenant_id=tenant_id,
            db=uow.session,
        )
