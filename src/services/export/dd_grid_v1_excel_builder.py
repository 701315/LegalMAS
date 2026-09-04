# src/services/export/dd_grid_v1_excel_builder.py

"""
DD Grid V1 Excel Builder.

Генерирует Excel-книгу из DDGridV1Data (новая модель) без зависимости от AggregatedResults.

Листы:
  1. "DD Grid"            — плоский список результатов из cell_values
  2. "DD Matrix"          — матричное представление из cell_values
  3. "Summary"            — KPI-сводка (прямой подсчёт из check_results)
  4. "By Doc Type"        — сводка по типам документов (группировка check_results)
  5. "Conditional Checks" — сработавшие условные проверки (мигрировано)
  6. "Extracted Values"   — структурированные значения (мигрировано)

Архитектурный принцип:
  - Никакой зависимости от aggregator.py
  - Чтение напрямую из DDGridV1Data + check_results
  - Детерминированный порядок строк/столбцов
  - Декларативная локализация через core.i18n.export_translator
"""

from __future__ import annotations

import io
from collections import defaultdict
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from core.config import settings
from core.i18n.export_translator import t_export
from core.logging import get_logger
from engine.playbook.models import CheckDefinition
from services.dd_grid_v1_builder import DDGridV1Data

logger = get_logger(__name__)


# ---------------------------------------------------------------------------
# Стили
# ---------------------------------------------------------------------------

_FILLS = {
    "critical": PatternFill(start_color="FF4444", end_color="FF4444", fill_type="solid"),
    "high": PatternFill(start_color="FF8C00", end_color="FF8C00", fill_type="solid"),
    "medium": PatternFill(start_color="FFD700", end_color="FFD700", fill_type="solid"),
    "pass": PatternFill(start_color="90EE90", end_color="90EE90", fill_type="solid"),
    "unclear": PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid"),
    "na": PatternFill(start_color="F5F5F5", end_color="F5F5F5", fill_type="solid"),
}
_FONT_WHITE = Font(color="FFFFFF", bold=True)
_FONT_BLACK = Font(color="000000")
_FONT_HEADER = Font(bold=True, size=11)
_FONT_TITLE = Font(bold=True, size=14)
_FONT_SMALL = Font(size=9, color="555555")
_THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)
_WRAP = Alignment(wrap_text=True, vertical="top")
_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


# ---------------------------------------------------------------------------
# Лист 1: DD Grid — плоский список
# ---------------------------------------------------------------------------

_BASE_COLUMN_SPECS = [
    ("columns.document", 35, "document_name"),
    ("columns.segment", 25, "segment_title"),
    ("columns.doc_type", 18, "doc_type"),
    ("columns.check", 30, "check_name"),
    ("columns.column", 25, "column_label"),
    ("columns.state", 15, "cell_state"),
    ("columns.value", 50, "display_value"),
    ("columns.page", 6, "location_page"),
    ("columns.clause", 10, "location_clause_ref"),
]


def _build_grid_sheet(wb: Workbook, grid_data: DDGridV1Data, locale: str = "en") -> None:
    """Лист 1: DD Grid — плоский список результатов."""
    ws = wb.active
    ws.title = t_export("sheets.grid", locale)
    ws.freeze_panes = "A2"

    # Заголовки
    for col_idx, (header_key, width, _) in enumerate(_BASE_COLUMN_SPECS, 1):
        cell = ws.cell(row=1, column=col_idx, value=t_export(header_key, locale))
        cell.font = _FONT_HEADER
        cell.border = _THIN_BORDER
        cell.alignment = _CENTER
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    # Собираем плоские строки из cell_values
    rows_data: list[dict] = []

    doc_meta_by_id = {dm.document_id: dm for dm in grid_data.documents_meta}

    drawer_by_doc_check: dict[tuple[str, str, str], Any] = {}
    for doc_id, segments in grid_data.drawer_contents.items():
        for seg_id, checks in segments.items():
            for check_id, drawer in checks.items():
                drawer_by_doc_check[(doc_id, seg_id, check_id)] = drawer

    column_registry = grid_data.column_registry

    for doc_id, segments in grid_data.cell_values.items():
        doc_meta = doc_meta_by_id.get(doc_id, {})
        doc_name = getattr(doc_meta, "filename", "")
        doc_type = getattr(doc_meta, "doc_type", "")

        for seg_id, columns in segments.items():
            segment_title = getattr(doc_meta, "segment_title", None) or seg_id

            for col_id, cell in columns.items():
                check_id = col_id.split("__")[0] if "__" in col_id else col_id

                drawer = drawer_by_doc_check.get((doc_id, seg_id, check_id))
                check_name = getattr(drawer, "check_name", check_id) if drawer else check_id

                column = column_registry.get(col_id)
                column_label = getattr(column, "label", col_id) if column else col_id

                rows_data.append({
                    "document_name": doc_name,
                    "segment_title": segment_title,
                    "doc_type": doc_type,
                    "check_name": check_name,
                    "column_label": column_label,
                    "cell_state": cell.cell_state,
                    "display_value": cell.display_value,
                    "location_page": cell.location_page,
                    "location_clause_ref": cell.location_clause_ref,
                    "severity": _infer_severity_from_state(cell.cell_state),
                })

    severity_order = {"critical": 0, "high": 1, "medium": 2, "unclear": 3, "pass": 4}
    rows_data.sort(key=lambda r: (
        severity_order.get(r["severity"], 9),
        r["document_name"].lower(),
        r["check_name"].lower(),
    ))

    for row_idx, row in enumerate(rows_data, 2):
        for col_idx, (_, _, field) in enumerate(_BASE_COLUMN_SPECS, 1):
            value = row.get(field, "")
            if value is None:
                value = ""

            cell = ws.cell(row=row_idx, column=col_idx, value=str(value))
            cell.border = _THIN_BORDER
            cell.alignment = _WRAP

        fill = _FILLS.get(row["severity"], _FILLS["unclear"])
        for col_idx in range(1, len(_BASE_COLUMN_SPECS) + 1):
            ws.cell(row=row_idx, column=col_idx).fill = fill
            if row["severity"] == "critical":
                ws.cell(row=row_idx, column=col_idx).font = _FONT_WHITE

    ws.auto_filter.ref = f"A1:{get_column_letter(len(_BASE_COLUMN_SPECS))}{len(rows_data) + 1}"


def _infer_severity_from_state(cell_state: str) -> str:
    mapping = {
        "error": "critical",
        "not_defined": "high",
        "not_executed": "unclear",
        "defined": "pass",
    }
    return mapping.get(cell_state, "unclear")


# ---------------------------------------------------------------------------
# Лист 2: DD Matrix — матричное представление
# ---------------------------------------------------------------------------


def _build_matrix_sheet(wb: Workbook, grid_data: DDGridV1Data, locale: str = "en") -> None:
    """Лист 2: DD Matrix — матрица документы × проверки."""
    ws = wb.create_sheet(t_export("sheets.matrix", locale))

    if not grid_data.cell_values:
        ws.cell(row=1, column=1, value=t_export("matrix.no_data", locale)).font = _FONT_TITLE
        return

    check_meta: dict[str, dict] = {}
    for doc_id, segments in grid_data.drawer_contents.items():
        for seg_id, checks in segments.items():
            for check_id, drawer in checks.items():
                if check_id not in check_meta:
                    check_meta[check_id] = {
                        "check_name": drawer.check_name,
                        "severity": _infer_severity_from_state(drawer.cell_state),
                    }

    if not check_meta:
        ws.cell(row=1, column=1, value=t_export("matrix.no_checks", locale)).font = _FONT_TITLE
        return

    doc_meta_by_id = {dm.document_id: dm for dm in grid_data.documents_meta}

    doc_stats: dict[str, dict] = defaultdict(lambda: {"critical": 0, "fail": 0})
    for doc_id, segments in grid_data.cell_values.items():
        for seg_id, columns in segments.items():
            for col_id, cell in columns.items():
                if cell.cell_state == "error":
                    doc_stats[doc_id]["critical"] += 1
                    doc_stats[doc_id]["fail"] += 1
                elif cell.cell_state == "not_defined":
                    doc_stats[doc_id]["fail"] += 1

    doc_order = sorted(
        grid_data.cell_values.keys(),
        key=lambda doc_id: (
            -doc_stats[doc_id]["critical"],
            -doc_stats[doc_id]["fail"],
            doc_meta_by_id.get(doc_id, {}).filename.lower() if hasattr(doc_meta_by_id.get(doc_id, {}), "filename") else "",
        )
    )

    fixed_cols = 3  # A–C (Document, Segment, Type)

    ws.cell(row=1, column=1, value=t_export("columns.document", locale)).font = _FONT_HEADER
    ws.cell(row=1, column=2, value=t_export("columns.segment", locale)).font = _FONT_HEADER
    ws.cell(row=1, column=3, value=t_export("columns.doc_type_short", locale)).font = _FONT_HEADER

    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 25
    ws.column_dimensions["C"].width = 16

    for col_idx in range(1, fixed_cols + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.alignment = _CENTER
        cell.border = _THIN_BORDER

    current_col = fixed_cols + 1
    check_ids = sorted(check_meta.keys(), key=lambda cid: check_meta[cid]["check_name"])

    for cid in check_ids:
        meta = check_meta[cid]
        severity = meta["severity"]
        check_name = meta["check_name"]
        cell = ws.cell(row=1, column=current_col, value=check_name[:30])
        cell.font = _FONT_HEADER
        cell.border = _THIN_BORDER
        cell.alignment = _CENTER
        cell.fill = _FILLS.get(severity, _FILLS["unclear"])
        if severity == "critical":
            cell.font = Font(bold=True, size=11, color="FFFFFF")
        ws.column_dimensions[get_column_letter(current_col)].width = 20
        current_col += 1

    ws.freeze_panes = ws.cell(row=2, column=fixed_cols + 1)

    state_labels = {
        "error": t_export("state.error", locale),
        "not_defined": t_export("state.not_defined", locale),
        "not_executed": t_export("state.not_executed", locale),
        "defined": t_export("state.defined", locale),
    }

    for row_idx, doc_id in enumerate(doc_order, 2):
        doc_meta = doc_meta_by_id.get(doc_id, {})
        doc_name = getattr(doc_meta, "filename", "")

        segments = grid_data.cell_values.get(doc_id, {})
        if not segments:
            continue

        for seg_idx, (seg_id, columns) in enumerate(segments.items()):
            segment_title = getattr(doc_meta, "segment_title", None) or seg_id
            doc_type = getattr(doc_meta, "doc_type", "")

            actual_row = row_idx + seg_idx

            ws.cell(row=actual_row, column=1, value=doc_name)
            ws.cell(row=actual_row, column=2, value=segment_title)
            ws.cell(row=actual_row, column=3, value=doc_type)

            for col in range(1, fixed_cols + 1):
                cell = ws.cell(row=actual_row, column=col)
                cell.border = _THIN_BORDER
                cell.alignment = _WRAP

            col = fixed_cols + 1
            for cid in check_ids:
                cell = ws.cell(row=actual_row, column=col)
                cell.border = _THIN_BORDER
                cell.alignment = _WRAP

                cell_value = columns.get(cid)
                if cell_value is None:
                    for full_col_id, val in columns.items():
                        if full_col_id.startswith(f"{cid}__"):
                            cell_value = val
                            break

                if cell_value is None:
                    cell.value = t_export("fallback.na", locale)
                    cell.fill = _FILLS["na"]
                    cell.font = _FONT_SMALL
                else:
                    state_label = state_labels.get(
                        cell_value.cell_state,
                        cell_value.cell_state,
                    )
                    cell.value = (
                        f"{state_label}\n{cell_value.display_value[:50]}"
                        if cell_value.display_value
                        else state_label
                    )
                    cell.fill = _FILLS.get(
                        _infer_severity_from_state(cell_value.cell_state),
                        _FILLS["unclear"],
                    )
                    cell.font = _FONT_WHITE if cell_value.cell_state == "error" else _FONT_BLACK
                col += 1

            ws.row_dimensions[actual_row].height = 42


# ---------------------------------------------------------------------------
# Лист 3: Summary — KPI
# ---------------------------------------------------------------------------


def _build_summary_sheet(
    wb: Workbook,
    check_results: list[dict[str, Any]],
    documents: dict[str, dict[str, Any]],
    deal_name: str,
    run_id: str,
    checks_mode: str,
    pipeline_strategy: str,
    locale: str = "en",
) -> None:
    """Лист 3: Summary — KPI-сводка (прямой подсчёт из check_results)."""
    ws = wb.create_sheet(t_export("sheets.summary", locale))
    ws.column_dimensions["A"].width = 50
    ws.column_dimensions["B"].width = 30

    ws.cell(
        row=1,
        column=1,
        value=t_export("summary.title", locale, deal_name=deal_name),
    ).font = _FONT_TITLE
    ws.merge_cells("A1:B1")

    meta_parts: list[str] = []
    if checks_mode:
        meta_parts.append(t_export("summary.checks_mode", locale, checks_mode=checks_mode))
    if pipeline_strategy:
        meta_parts.append(
            t_export("summary.pipeline_strategy", locale, pipeline_strategy=pipeline_strategy)
        )
    if meta_parts:
        meta_text = " | ".join(meta_parts)
        cell = ws.cell(row=2, column=1, value=meta_text)
        cell.font = _FONT_SMALL
        cell.alignment = _WRAP
        ws.merge_cells("A2:B2")

    severity_counts = {"critical": 0, "high": 0, "medium": 0, "pass": 0, "unclear": 0}
    for cr in check_results:
        result = cr.get("result", "unclear")
        severity = cr.get("severity", "medium")

        if result == "fail":
            if severity in severity_counts:
                severity_counts[severity] += 1
        elif result == "pass":
            severity_counts["pass"] += 1
        else:
            severity_counts["unclear"] += 1

    total_documents = len(documents)
    total_checks = len(check_results)

    conditional_total = sum(1 for cr in check_results if cr.get("conditional_check_triggered"))
    conditional_fail = sum(
        1 for cr in check_results
        if cr.get("conditional_check_triggered") and cr.get("result") == "fail"
    )

    extracted_total = 0
    for cr in check_results:
        structured_data = cr.get("structured_data") or {}
        for field_value in structured_data.values():
            if isinstance(field_value, dict) and "value" in field_value:
                raw_value = field_value.get("value")
            else:
                raw_value = field_value
            if raw_value and str(raw_value) not in ("not_found", "not_defined", "None", ""):
                extracted_total += 1

    kpi_data = [
        (t_export("summary.total_documents", locale), total_documents),
        (t_export("summary.total_checks", locale), total_checks),
        ("", ""),
        ("🔴 Critical", severity_counts["critical"]),
        ("🟠 High", severity_counts["high"]),
        ("🟡 Medium", severity_counts["medium"]),
        ("🟢 Pass", severity_counts["pass"]),
        ("⚪ Unclear", severity_counts["unclear"]),
        ("", ""),
        (t_export("summary.conditional_total", locale), conditional_total),
        (t_export("summary.conditional_fail", locale), conditional_fail),
        (t_export("summary.extracted_values", locale), extracted_total),
    ]

    start_row = 4
    for row, (label, value) in enumerate(kpi_data, start_row):
        label_cell = ws.cell(row=row, column=1, value=label)
        label_cell.font = Font(bold=bool(label))
        val_cell = ws.cell(row=row, column=2, value=value)
        val_cell.alignment = Alignment(horizontal="right")

        if "Critical" in str(label):
            val_cell.fill = _FILLS["critical"]
            val_cell.font = _FONT_WHITE
        elif "High" in str(label):
            val_cell.fill = _FILLS["high"]
        elif "Medium" in str(label):
            val_cell.fill = _FILLS["medium"]
        elif "Pass" in str(label):
            val_cell.fill = _FILLS["pass"]
        elif "Unclear" in str(label):
            val_cell.fill = _FILLS["unclear"]


# ---------------------------------------------------------------------------
# Лист 4: By Doc Type
# ---------------------------------------------------------------------------


def _build_doc_type_sheet(
    wb: Workbook,
    check_results: list[dict[str, Any]],
    documents: dict[str, dict[str, Any]],
    locale: str = "en",
) -> None:
    """Лист 4: By Doc Type — сводка по типам документов."""
    ws = wb.create_sheet(t_export("sheets.by_doc_type", locale))
    ws.freeze_panes = "A2"

    headers = [
        (t_export("columns.doc_type", locale), 28),
        (t_export("columns.documents_count", locale), 12),
        (t_export("columns.checks_count", locale), 12),
        ("🔴 Critical", 12),
        ("🟠 High", 12),
        ("🟡 Medium", 12),
        ("🟢 Pass", 12),
        ("⚪ Unclear", 12),
        (t_export("columns.conditional_fail", locale), 16),
    ]

    for col_idx, (header, width) in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = _FONT_HEADER
        cell.border = _THIN_BORDER
        cell.alignment = _CENTER
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    by_type: dict[str, dict] = defaultdict(lambda: {
        "documents": set(),
        "checks": set(),
        "critical": 0,
        "high": 0,
        "medium": 0,
        "pass": 0,
        "unclear": 0,
        "conditional_fail": 0,
    })

    for cr in check_results:
        doc_id = cr.get("document_id", "")
        doc_type = documents.get(doc_id, {}).get("doc_type", t_export("fallback.generic", locale))

        by_type[doc_type]["documents"].add(doc_id)
        by_type[doc_type]["checks"].add(cr.get("check_id", ""))

        result = cr.get("result", "unclear")
        severity = cr.get("severity", "medium")

        if result == "fail":
            if severity in ("critical", "high", "medium"):
                by_type[doc_type][severity] += 1
        elif result == "pass":
            by_type[doc_type]["pass"] += 1
        else:
            by_type[doc_type]["unclear"] += 1

        if cr.get("conditional_check_triggered") and result == "fail":
            by_type[doc_type]["conditional_fail"] += 1

    sorted_summaries = sorted(
        by_type.items(),
        key=lambda x: -(x[1]["critical"] + x[1]["high"])
    )

    for row_idx, (doc_type, stats) in enumerate(sorted_summaries, 2):
        values = [
            doc_type,
            len(stats["documents"]),
            len(stats["checks"]),
            stats["critical"],
            stats["high"],
            stats["medium"],
            stats["pass"],
            stats["unclear"],
            stats["conditional_fail"],
        ]

        for col_idx, value in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.border = _THIN_BORDER
            cell.alignment = _WRAP

        if stats["critical"] > 0:
            ws.cell(row=row_idx, column=4).fill = _FILLS["critical"]
            ws.cell(row=row_idx, column=4).font = _FONT_WHITE
        if stats["high"] > 0:
            ws.cell(row=row_idx, column=5).fill = _FILLS["high"]


# ---------------------------------------------------------------------------
# Лист 5: Conditional Checks
# ---------------------------------------------------------------------------


def _build_conditional_sheet(
    wb: Workbook,
    check_results: list[dict[str, Any]],
    documents: dict[str, dict[str, Any]],
    locale: str = "en",
) -> None:
    """Лист 5: Conditional Checks — сработавшие условные проверки."""
    ws = wb.create_sheet(t_export("sheets.conditional", locale))
    ws.freeze_panes = "A2"

    headers = [
        (t_export("columns.document", locale), 35),
        (t_export("columns.doc_type_short", locale), 15),
        (t_export("columns.trigger", locale), 22),
        (t_export("columns.check", locale), 30),
        (t_export("columns.result", locale), 12),
        (t_export("columns.severity", locale), 12),
        (t_export("columns.evidence", locale), 45),
        (t_export("columns.page", locale), 6),
        (t_export("columns.confidence", locale), 12),
    ]

    for col, (h, w) in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = _FONT_HEADER
        cell.border = _THIN_BORDER
        cell.alignment = _CENTER
        ws.column_dimensions[get_column_letter(col)].width = w

    conditional_rows = [
        cr for cr in check_results
        if cr.get("conditional_check_triggered")
    ]

    for row_idx, cr in enumerate(conditional_rows, 2):
        doc_id = cr.get("document_id", "")
        doc_info = documents.get(doc_id, {})

        evidence = cr.get("display_text") or cr.get("value_text", "")
        evidence_snippet = evidence[:300] if evidence else ""

        values = [
            doc_info.get("name", t_export("fallback.unknown", locale)),
            doc_info.get("doc_type", t_export("fallback.generic", locale)),
            cr.get("trigger_name", ""),
            cr.get("check_name", ""),
            cr.get("result", "unclear"),
            cr.get("severity", "medium"),
            evidence_snippet,
            cr.get("page"),
            f"{cr.get('confidence', 0.0):.0%}",
        ]

        for col_idx, value in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=str(value) if value else "")
            cell.border = _THIN_BORDER
            cell.alignment = _WRAP

        severity = cr.get("severity", "medium")
        fill = _FILLS.get(severity, _FILLS["unclear"])
        for col_idx in range(1, len(headers) + 1):
            ws.cell(row=row_idx, column=col_idx).fill = fill
            if severity == "critical":
                ws.cell(row=row_idx, column=col_idx).font = _FONT_WHITE


# ---------------------------------------------------------------------------
# Лист 6: Extracted Values
# ---------------------------------------------------------------------------


def _infer_field_format(field_name: str, value: str) -> str:
    name_lower = field_name.lower()
    if any(x in name_lower for x in ("amount", "price", "sum", "cost")):
        return "money"
    if any(x in name_lower for x in ("date", "expiry", "term")):
        return "date"
    if any(x in name_lower for x in ("party", "parties", "entity")):
        return "entity_name"
    return "text"


def _build_extracted_values_sheet(
    wb: Workbook,
    check_results: list[dict[str, Any]],
    documents: dict[str, dict[str, Any]],
    locale: str = "en",
) -> None:
    """Лист 6: Extracted Values — структурированные значения."""
    ws = wb.create_sheet(t_export("sheets.extracted", locale))
    ws.freeze_panes = "A2"

    headers = [
        (t_export("columns.document", locale), 35),
        (t_export("columns.doc_type_short", locale), 15),
        (t_export("columns.check", locale), 30),
        (t_export("columns.field", locale), 20),
        (t_export("columns.value", locale), 35),
        (t_export("columns.format", locale), 15),
        (t_export("columns.page", locale), 6),
        (t_export("columns.clause", locale), 10),
    ]
    fields = [
        "document_name",
        "doc_type",
        "check_name",
        "field_name",
        "field_value",
        "field_format",
        "page",
        "clause_uid",
    ]

    for col, (h, w) in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = _FONT_HEADER
        cell.border = _THIN_BORDER
        cell.alignment = _CENTER
        ws.column_dimensions[get_column_letter(col)].width = w

    extracted_rows: list[dict] = []

    for cr in check_results:
        doc_id = cr.get("document_id", "")
        doc_info = documents.get(doc_id, {})
        structured_data = cr.get("structured_data") or {}

        for field_name, field_value in structured_data.items():
            if isinstance(field_value, dict) and "value" in field_value:
                raw_display_value = field_value.get("value")
            else:
                raw_display_value = field_value

            if raw_display_value is None or str(raw_display_value) in ("not_found", "not_defined", "None", ""):
                continue

            display_value = str(raw_display_value)

            extracted_rows.append({
                "document_name": doc_info.get("name", t_export("fallback.unknown", locale)),
                "doc_type": doc_info.get("doc_type", t_export("fallback.generic", locale)),
                "check_name": cr.get("check_name", ""),
                "field_name": field_name,
                "field_value": display_value,
                "field_format": _infer_field_format(field_name, display_value),
                "page": cr.get("page"),
                "clause_uid": cr.get("clause_uid", ""),
            })

    for row_idx, row in enumerate(extracted_rows, 2):
        for col_idx, field in enumerate(fields, 1):
            value = row.get(field, "")
            if field == "page" and value is None:
                value = ""
            cell = ws.cell(row=row_idx, column=col_idx, value=str(value) if value else "")
            cell.border = _THIN_BORDER
            cell.alignment = _WRAP


# ---------------------------------------------------------------------------
# Точка входа
# ---------------------------------------------------------------------------


def build_dd_grid_v1_excel(
    grid_data: DDGridV1Data,
    check_results: list[dict[str, Any]],
    documents: dict[str, dict[str, Any]],
    checks_by_id: dict[str, CheckDefinition],
    deal_name: str,
    run_id: str,
    checks_mode: str,
    pipeline_strategy: str,
    locale: str = "en",
) -> bytes:
    """
    Построить полную Excel-книгу из DDGridV1Data.

    Args:
        grid_data: DDGridV1Data из build_dd_grid_v1()
        check_results: Сырые результаты проверок из БД
        documents: Метаданные документов {doc_id: {...}}
        checks_by_id: Словарь {check_id: CheckDefinition}
        deal_name: Название сделки для шапки
        run_id: ID Run для логирования
        checks_mode: Режим проверок (для шапки)
        pipeline_strategy: Стратегия пайплайна (для шапки)
        locale: presentation locale for export labels (en canonical default)

    Returns:
        Excel file as bytes
    """
    logger.info(
        "Building DD Grid V1 Excel",
        run_id=run_id,
        deal_name=deal_name,
        documents=len(grid_data.documents_meta),
        checks=len(grid_data.column_registry),
        locale=locale,
    )

    wb = Workbook()

    _build_grid_sheet(wb, grid_data, locale=locale)
    _build_matrix_sheet(wb, grid_data, locale=locale)
    _build_summary_sheet(
        wb,
        check_results,
        documents,
        deal_name,
        run_id,
        checks_mode,
        pipeline_strategy,
        locale=locale,
    )
    _build_doc_type_sheet(
        wb,
        check_results,
        documents,
        locale=locale,
    )
    _build_conditional_sheet(
        wb,
        check_results,
        documents,
        locale=locale,
    )
    _build_extracted_values_sheet(
        wb,
        check_results,
        documents,
        locale=locale,
    )

    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    logger.info(
        "DD Grid V1 Excel built",
        run_id=run_id,
        sheets=len(wb.sheetnames),
        locale=locale,
    )

    return output.read()


    