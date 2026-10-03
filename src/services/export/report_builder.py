# src/services/export/report_builder.py

"""
Word report builder based on legal due diligence analysis results (DD Report).

What it does:
- Generates a multi-page .docx report for a single Run;
- Adds a title page with deal details and analysis context;
- Consolidates key metrics and risks into user-friendly sections for lawyers.

Report structure:
- Title page:
  - Report and deal names;
  - Playbook profile, version, Run ID;
  - Checks mode and pipeline strategy.
- 1. Executive Summary — brief summary of severity counts and overall analysis scope.
- 2. Critical Findings — detailed list of critical/high risks.
- 3. Analysis Quality — quarantine, degraded documents, evidence validity.
- 4. Conditional Checks — triggered conditional checks (if any).
- 5. Extracted Values — table of structured data (if any).
- Appendix A: Request List — DD requests with addressee and document bindings.
- Reproducibility footer — concise footprint of analysis context (profile, playbook, Run, env).

Features:
- Tailored for lawyers: minimal technical jargon, maximum readability;
- All key Run parameters explicitly baked into the cover page and footer for auditing;
- Declarative localization via core.i18n.export_translator.

Architectural principle:
- Zero dependency on aggregator.py
- Direct reading from check_results + documents + checks_by_id
- Deterministic section ordering and sorting
"""

from __future__ import annotations

import io
from datetime import datetime
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from core.i18n.export_translator import t_export
from core.logging import get_logger
from engine.playbook.models import CheckDefinition
from services.dd_request_rules import should_trigger_dd_request
from services.dd_request_composer import compose_dd_request

from engine.checks.result_semantics import get_display_clause_ref
logger = get_logger(__name__)

_SEVERITY_COLORS = {
    "critical": RGBColor(0xFF, 0x44, 0x44),
    "high": RGBColor(0xFF, 0x8C, 0x00),
    "medium": RGBColor(0xFF, 0xD7, 0x00),
    "pass": RGBColor(0x00, 0x80, 0x00),
    "unclear": RGBColor(0x80, 0x80, 0x80),
}


# ---------------------------------------------------------------------------
# Helpers — direct computation of metrics from check_results
# ---------------------------------------------------------------------------


def _count_severity(check_results: list[dict[str, Any]]) -> dict[str, int]:
    """Compute severity counts directly from check_results."""
    counts = {"critical": 0, "high": 0, "medium": 0, "pass": 0, "unclear": 0}
    for cr in check_results:
        result = cr.get("result", "unclear")
        severity = cr.get("severity", "medium")

        if result == "fail":
            if severity in counts:
                counts[severity] += 1
        elif result == "pass":
            counts["pass"] += 1
        else:
            counts["unclear"] += 1
    return counts


def _count_conditional(check_results: list[dict[str, Any]]) -> dict[str, int]:
    """Compute conditional checks statistics."""
    total = sum(1 for cr in check_results if cr.get("conditional_check_triggered"))
    triggered_fail = sum(
        1 for cr in check_results
        if cr.get("conditional_check_triggered") and cr.get("result") == "fail"
    )
    return {"total": total, "triggered_fail": triggered_fail}


def _get_top_risks(
    check_results: list[dict[str, Any]],
    documents: dict[str, dict[str, Any]],
    limit: int = 15,
    locale: str = "en",
) -> list[dict]:
    """Get top risks (critical/high) from check_results."""
    severity_order = {"critical": 0, "high": 1, "medium": 2}

    fail_rows = [
        cr for cr in check_results
        if cr.get("result") == "fail" and cr.get("severity") in ("critical", "high")
    ]

    fail_rows.sort(key=lambda r: (
        severity_order.get(r.get("severity", "medium"), 9),
        -r.get("confidence", 0.0),
    ))

    risks = []
    for cr in fail_rows[:limit]:
        doc_id = cr.get("document_id", "")
        doc_info = documents.get(doc_id, {})

        risks.append({
            "check_name": cr.get("check_name", ""),
            "severity": cr.get("severity", "medium"),
            "document_name": doc_info.get("name", t_export("fallback.unknown", locale)),
            "doc_type": doc_info.get("doc_type", t_export("fallback.generic", locale)),
            "clause_uid": get_display_clause_ref(
                cr.get("result", ""), cr.get("clause_uid", "")
            ),
            "page": cr.get("page"),
            "evidence_snippet": (cr.get("display_text") or cr.get("value_text", ""))[:300],
            "structured_data": cr.get("structured_data") or {},
            "conditional_check_triggered": cr.get("conditional_check_triggered", False),
            "trigger_name": cr.get("trigger_name", ""),
        })

    return risks


def _count_quality_metrics(
    check_results: list[dict[str, Any]],
    documents: dict[str, dict[str, Any]],
) -> dict[str, Any]:
    """Compute analysis quality metrics."""
    quarantine_count = sum(
        1 for doc in documents.values()
        if doc.get("status") == "quarantine"
    )

    degraded_docs = set()
    for cr in check_results:
        quality_flags = cr.get("quality_flags") or {}
        if isinstance(quality_flags, dict) and quality_flags.get("structure_degraded"):
            degraded_docs.add(cr.get("document_id", ""))

    fail_high_crit = [
        cr for cr in check_results
        if cr.get("result") == "fail" and cr.get("severity") in ("critical", "high")
    ]

    evidence_valid = sum(
        1 for cr in fail_high_crit
        if (cr.get("display_text") or cr.get("value_text") or "").strip()
    )

    evidence_validity_rate = (
        evidence_valid / len(fail_high_crit) if fail_high_crit else 1.0
    )

    return {
        "quarantine_count": quarantine_count,
        "degraded_count": len(degraded_docs),
        "evidence_validity_rate": evidence_validity_rate,
    }


def _get_conditional_rows(
    check_results: list[dict[str, Any]],
    documents: dict[str, dict[str, Any]],
    limit: int = 20,
    locale: str = "en",
) -> list[dict]:
    """Get triggered conditional checks."""
    rows = []
    for cr in check_results:
        if not cr.get("conditional_check_triggered"):
            continue

        doc_id = cr.get("document_id", "")
        doc_info = documents.get(doc_id, {})

        rows.append({
            "document_name": doc_info.get("name", t_export("fallback.unknown", locale)),
            "doc_type": doc_info.get("doc_type", t_export("fallback.generic", locale)),
            "trigger_name": cr.get("trigger_name", ""),
            "check_name": cr.get("check_name", ""),
            "result": cr.get("result", "unclear"),
            "severity": cr.get("severity", "medium"),
            "evidence_snippet": (cr.get("display_text") or cr.get("value_text", ""))[:150],
        })

        if len(rows) >= limit:
            break

    return rows


def _get_extracted_values(
    check_results: list[dict[str, Any]],
    documents: dict[str, dict[str, Any]],
    limit: int = 30,
    locale: str = "en",
) -> list[dict]:
    """Extract structured values from check_results."""
    rows = []

    for cr in check_results:
        doc_id = cr.get("document_id", "")
        doc_info = documents.get(doc_id, {})
        structured_data = cr.get("structured_data") or {}

        for field_name, field_value in structured_data.items():
            if isinstance(field_value, dict) and "value" in field_value:
                raw_display_value = field_value.get("value")
            else:
                raw_display_value = field_value

            if raw_display_value is None or str(raw_display_value) in ("not_found", "not_defined", "None", ""):
                continue

            display_value = str(raw_display_value)

            rows.append({
                "document_name": doc_info.get("name", t_export("fallback.unknown", locale)),
                "doc_type": doc_info.get("doc_type", t_export("fallback.generic", locale)),
                "field_name": field_name,
                "field_value": display_value,
                "clause_uid": get_display_clause_ref(
                    cr.get("result", ""), cr.get("clause_uid", "")
                ),
                "page": cr.get("page"),
            })

            if len(rows) >= limit:
                break

        if len(rows) >= limit:
            break

    return rows


def _build_dd_requests_for_report(
    check_results: list[dict[str, Any]],
    documents: dict[str, dict[str, Any]],
    checks_by_id: dict[str, CheckDefinition],
    locale: str = "en",
) -> list[dict]:
    """
    Build DD requests for the report.
    """
    from services.export.dd_request_loader import get_dd_request_loader

    loader = get_dd_request_loader()

    results_by_check: dict[str, list[dict]] = {}
    for cr in check_results:
        check_id = cr.get("check_id", "")
        if not check_id:
            continue
        results_by_check.setdefault(check_id, []).append(cr)

    requests = []

    for check_id, cr_list in results_by_check.items():
        check_def = checks_by_id.get(check_id)
        if not check_def:
            continue

        # B7.9.2 / Variant B: DD-request trigger via rules engine.
        override_results = getattr(check_def, "dd_request_when_result", None)

        triggered_docs: list[dict] = []
        trigger_cr: dict | None = None
        for cr in cr_list:
            if not should_trigger_dd_request(
                check_result=cr,
                override_results=override_results,
            ):
                continue
            if trigger_cr is None:
                trigger_cr = cr
            doc_id = cr.get("document_id", "")
            structured_data = cr.get("structured_data") or {}
            triggered_docs.append({
                "document_id": doc_id,
                "document_name": documents.get(doc_id, {}).get("name", t_export("fallback.unknown", locale)),
                "doc_type": documents.get(doc_id, {}).get("doc_type", t_export("fallback.generic", locale)),
                "structured_data": structured_data,
            })

        if not triggered_docs:
            continue

        doc_names = [d["document_name"] for d in triggered_docs]
        text, addressee = compose_dd_request(
            check_result=trigger_cr,
            check_def=check_def,
            doc_names=doc_names,
            loader=loader,
            locale=locale,
        )

        # required_fields — только для structured_values (best-effort).
        required_fields: list[str] = []
        extraction = check_def.extraction
        if extraction is not None and getattr(extraction, "enabled", False):
            required_fields = [
                f.name for f in extraction.fields
                if getattr(f, "required", False)
            ]

        structured_values = []
        for doc in triggered_docs:
            for field_name in required_fields:
                field_value = doc["structured_data"].get(field_name)
                if field_value:
                    if isinstance(field_value, dict) and "value" in field_value:
                        display_value = str(field_value.get("value", ""))
                    else:
                        display_value = str(field_value)
                    if display_value and display_value not in ("not_found", "not_defined", "None", ""):
                        structured_values.append({
                            "field": field_name,
                            "value": display_value,
                        })

        requests.append({
            "priority": check_def.severity,
            "check_name": check_def.name,
            "addressee": addressee,
            "text": text,
            "document_count": len(triggered_docs),
            "document_names": doc_names[:5],
            "structured_values": structured_values[:5],
        })

    severity_order = {"critical": 0, "high": 1, "medium": 2}
    requests.sort(key=lambda r: severity_order.get(r["priority"], 9))

    return requests


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------


def build_report_docx(
    check_results: list[dict[str, Any]],
    documents: dict[str, dict[str, Any]],
    checks_by_id: dict[str, CheckDefinition],
    deal_name: str,
    run_id: str,
    playbook_version: str = "1.0.0",
    playbook_profile: str = "M&A Russia DD",
    environment_hash: str = "",
    checks_mode: str = "",
    pipeline_strategy: str = "",
    locale: str = "en",
) -> bytes:
    """
    Generate DD Report as Word document based on single Run results.
    """
    doc = Document()

    _add_title_page(
        doc,
        deal_name=deal_name,
        run_id=run_id,
        pb_version=playbook_version,
        pb_profile=playbook_profile,
        checks_mode=checks_mode,
        pipeline_strategy=pipeline_strategy,
        locale=locale,
    )
    _add_executive_summary(doc, check_results, documents, locale=locale)
    _add_critical_findings(doc, check_results, documents, locale=locale)
    _add_quality_section(doc, check_results, documents, locale=locale)

    conditional_rows = _get_conditional_rows(check_results, documents, locale=locale)
    if conditional_rows:
        _add_conditional_checks_section(doc, conditional_rows, locale=locale)

    extracted_values = _get_extracted_values(check_results, documents, locale=locale)
    if extracted_values:
        _add_extracted_values_section(doc, extracted_values, locale=locale)

    dd_requests = _build_dd_requests_for_report(check_results, documents, checks_by_id, locale=locale)
    _add_request_list(doc, dd_requests, locale=locale)

    _add_footer(
        doc,
        run_id=run_id,
        pb_version=playbook_version,
        pb_profile=playbook_profile,
        env_hash=environment_hash,
    )

    output = io.BytesIO()
    doc.save(output)
    output.seek(0)
    logger.info("DD Report Word built (V1 path)", locale=locale)
    return output.read()


# ---------------------------------------------------------------------------
# Report sections
# ---------------------------------------------------------------------------


def _add_title_page(
    doc: Document,
    *,
    deal_name: str,
    run_id: str,
    pb_version: str,
    pb_profile: str,
    checks_mode: str = "",
    pipeline_strategy: str = "",
    locale: str = "en",
) -> None:
    doc.add_paragraph("")
    doc.add_paragraph("")
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(t_export("report.title", locale))
    run.bold = True
    run.font.size = Pt(18)

    doc.add_paragraph("")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run(deal_name).font.size = Pt(14)

    doc.add_paragraph("")
    info = doc.add_paragraph()
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    info.add_run(f"{t_export('report.profile', locale, profile=pb_profile)}\n").font.size = Pt(11)
    info.add_run(f"{t_export('report.date', locale, date=datetime.now().strftime('%d.%m.%Y'))}\n").font.size = Pt(11)
    info.add_run(f"{t_export('report.playbook', locale, version=pb_version)}\n").font.size = Pt(11)
    info.add_run(f"{t_export('report.run_id', locale, run_id=run_id[:8])}\n").font.size = Pt(9)

    if checks_mode:
        info.add_run(f"{t_export('report.checks_mode', locale, mode=checks_mode)}\n").font.size = Pt(10)
    if pipeline_strategy:
        info.add_run(f"{t_export('report.pipeline_strategy', locale, strategy=pipeline_strategy)}\n").font.size = Pt(10)

    doc.add_page_break()


def _add_executive_summary(
    doc: Document,
    check_results: list[dict[str, Any]],
    documents: dict[str, dict[str, Any]],
    locale: str = "en",
) -> None:
    doc.add_heading(t_export("report.sec1_title", locale), level=1)

    counts = _count_severity(check_results)
    conditional_stats = _count_conditional(check_results)

    doc.add_paragraph(
        t_export(
            "report.summary_scope",
            locale,
            doc_count=len(documents),
            check_count=len(check_results),
        )
    )

    if conditional_stats["total"] > 0:
        doc.add_paragraph(
            t_export(
                "report.summary_conditional",
                locale,
                total=conditional_stats["total"],
                fail_count=conditional_stats["triggered_fail"],
            )
        )

    table = doc.add_table(rows=6, cols=2)
    table.style = "Light Grid Accent 1"
    cells = [
        ("🔴 Critical", str(counts["critical"])),
        ("🟠 High", str(counts["high"])),
        ("🟡 Medium", str(counts["medium"])),
        ("🟢 Pass", str(counts["pass"])),
        ("⚪ Unclear", str(counts["unclear"])),
        (t_export("report.total", locale), str(sum(counts.values()))),
    ]
    for i, (label, value) in enumerate(cells):
        table.cell(i, 0).text = label
        table.cell(i, 1).text = value

    doc.add_paragraph("")
    if counts["critical"] > 0:
        assessment = t_export("report.assessment_critical", locale, count=counts["critical"])
    elif counts["high"] > 0:
        assessment = t_export("report.assessment_high", locale, count=counts["high"])
    else:
        assessment = t_export("report.assessment_none", locale)

    p = doc.add_paragraph(assessment)
    if counts["critical"] > 0:
        p.runs[0].bold = True


def _add_critical_findings(
    doc: Document,
    check_results: list[dict[str, Any]],
    documents: dict[str, dict[str, Any]],
    locale: str = "en",
) -> None:
    doc.add_heading(t_export("report.sec2_title", locale), level=1)
    critical_high = _get_top_risks(check_results, documents, limit=15, locale=locale)

    if not critical_high:
        doc.add_paragraph(t_export("report.assessment_none", locale))
        return

    for i, risk in enumerate(critical_high, 1):
        doc.add_heading(
            f"2.{i}. {risk['check_name']} ({risk['severity'].upper()})",
            level=2,
        )
        doc.add_paragraph(
            t_export(
                "report.document_label",
                locale,
                doc_name=risk["document_name"],
                doc_type=risk["doc_type"],
            )
        )
        if risk["clause_uid"]:
            if risk["page"]:
                doc.add_paragraph(
                    t_export("report.clause_page", locale, clause=risk["clause_uid"], page=risk["page"])
                )
            else:
                doc.add_paragraph(
                    t_export("report.clause_only", locale, clause=risk["clause_uid"])
                )
        elif risk["page"]:
            doc.add_paragraph(t_export("report.page_only", locale, page=risk["page"]))

        if risk["evidence_snippet"]:
            quote = doc.add_paragraph()
            run = quote.add_run(f"«{risk['evidence_snippet']}»")
            run.italic = True

        if risk["structured_data"]:
            sd_text = "; ".join(f"{k}: {v}" for k, v in risk["structured_data"].items())
            doc.add_paragraph(t_export("report.extracted_data", locale, data=sd_text))

        if risk["conditional_check_triggered"] and risk["trigger_name"]:
            doc.add_paragraph(
                t_export("report.triggered_by", locale, trigger_name=risk["trigger_name"])
            )
        doc.add_paragraph("")


def _add_quality_section(
    doc: Document,
    check_results: list[dict[str, Any]],
    documents: dict[str, dict[str, Any]],
    locale: str = "en",
) -> None:
    doc.add_heading(t_export("report.sec3_title", locale), level=1)
    metrics = _count_quality_metrics(check_results, documents)

    items = [
        t_export("report.quarantine_docs", locale, count=metrics["quarantine_count"]),
        t_export("report.degraded_docs", locale, count=metrics["degraded_count"]),
        t_export("report.evidence_validity_rate", locale, rate=f"{metrics['evidence_validity_rate']:.0%}"),
    ]
    for item in items:
        doc.add_paragraph(item, style="List Bullet")

    if metrics["quarantine_count"] > 0:
        doc.add_paragraph(t_export("report.quarantine_warning", locale))


def _add_conditional_checks_section(
    doc: Document,
    conditional_rows: list[dict],
    locale: str = "en",
) -> None:
    """Section for triggered conditional checks."""
    doc.add_heading(t_export("report.sec4_title", locale), level=1)
    doc.add_paragraph(t_export("report.conditional_desc", locale))

    fail_cond = [r for r in conditional_rows if r["result"] == "fail"]
    if not fail_cond:
        doc.add_paragraph(t_export("report.no_conditional_violations", locale))
        return

    table = doc.add_table(rows=1, cols=5)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    hdr[0].text = t_export("columns.document", locale)
    hdr[1].text = t_export("columns.trigger", locale)
    hdr[2].text = t_export("columns.check", locale)
    hdr[3].text = t_export("columns.severity", locale)
    hdr[4].text = t_export("columns.evidence", locale)

    for cr in fail_cond[:20]:
        row_cells = table.add_row().cells
        row_cells[0].text = cr["document_name"]
        row_cells[1].text = cr["trigger_name"]
        row_cells[2].text = cr["check_name"]
        row_cells[3].text = cr["severity"].upper()
        row_cells[4].text = cr["evidence_snippet"]


def _add_extracted_values_section(
    doc: Document,
    extracted_values: list[dict],
    locale: str = "en",
) -> None:
    """Section for structured extracted values (amounts, dates, etc.)."""
    doc.add_heading(t_export("report.sec5_title", locale), level=1)
    doc.add_paragraph(t_export("report.extracted_desc", locale))

    table = doc.add_table(rows=1, cols=5)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    hdr[0].text = t_export("columns.document", locale)
    hdr[1].text = t_export("columns.doc_type_short", locale)
    hdr[2].text = t_export("columns.field", locale)
    hdr[3].text = t_export("columns.value", locale)
    hdr[4].text = t_export("columns.clause", locale)

    for sv in extracted_values[:30]:
        row_cells = table.add_row().cells
        row_cells[0].text = sv["document_name"]
        row_cells[1].text = sv["doc_type"]
        row_cells[2].text = sv["field_name"]
        row_cells[3].text = sv["field_value"]
        row_cells[4].text = sv["clause_uid"] or (str(sv["page"]) if sv["page"] else "")


def _add_request_list(
    doc: Document,
    dd_requests: list[dict],
    locale: str = "en",
) -> None:
    doc.add_heading(t_export("report.appA_title", locale), level=1)
    if not dd_requests:
        doc.add_paragraph(t_export("report.no_requests", locale))
        return

    for i, req in enumerate(dd_requests, 1):
        doc.add_heading(
            f"A.{i}. [{req['priority'].upper()}] {req['check_name']}",
            level=2,
        )
        doc.add_paragraph(t_export("report.addressee", locale, addressee=req["addressee"]))
        doc.add_paragraph(req["text"])
        doc.add_paragraph(
            t_export(
                "report.documents_count",
                locale,
                count=req["document_count"],
                names=", ".join(req["document_names"][:5]),
            )
        )
        if req["structured_values"]:
            sv_text = "; ".join(f"{sv['field']}={sv['value']}" for sv in req["structured_values"][:5])
            doc.add_paragraph(t_export("report.linked_values", locale, values=sv_text))
        doc.add_paragraph("")


def _add_footer(
    doc: Document,
    *,
    run_id: str,
    pb_version: str,
    pb_profile: str,
    env_hash: str,
) -> None:
    doc.add_paragraph("")
    doc.add_paragraph("—" * 60)
    footer_text = (
        f"Reproducibility: Profile={pb_profile} | "
        f"Playbook v{pb_version} | "
        f"Run {run_id[:12]} | "
        f"Env {env_hash[:12] if env_hash else 'dev'} | "
        f"Generated {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
    )
    p = doc.add_paragraph(footer_text)
    p.runs[0].font.size = Pt(8)
    p.runs[0].font.color.rgb = RGBColor(0x80, 0x80, 0x80)


    